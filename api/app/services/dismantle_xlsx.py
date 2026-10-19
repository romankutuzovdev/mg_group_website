"""Build Excel workbook for a deal dismantle map (UK or USA layout)."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models.cabinet import DismantleMapOut


def _cell_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def build_dismantle_xlsx(data: DismantleMapOut) -> bytes:
    if (data.variant or "").lower() == "usa":
        return _build_usa_xlsx(data)
    return _build_uk_xlsx(data)


def _section_order(data: DismantleMapOut) -> list[str]:
    order = list(data.sections or [])
    if order:
        return order
    seen: list[str] = []
    for c in data.cells:
        if c.section and c.section not in seen:
            seen.append(c.section)
    return seen


def _build_uk_xlsx(data: DismantleMapOut) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Карта разбора"

    thin = Border(
        left=Side(style="thin", color="D4D4D4"),
        right=Side(style="thin", color="D4D4D4"),
        top=Side(style="thin", color="D4D4D4"),
        bottom=Side(style="thin", color="D4D4D4"),
    )
    header_fill = PatternFill("solid", fgColor="0D3F10")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    section_fill = PatternFill("solid", fgColor="E8F5E9")
    section_font = Font(bold=True, color="0D3F10", size=11)
    title_font = Font(bold=True, size=14, color="0D3F10")
    label_font = Font(size=10, color="666666")

    ws.merge_cells("A1:E1")
    ws["A1"] = _cell_value(data.title or "Карта разбора")
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(vertical="center")

    ws.merge_cells("A2:E2")
    hint = data.vehicle_header_hint or "марка авто, модель, Lot#"
    ws["A2"] = _cell_value(f"{hint}: {data.vehicle_label or '—'}")
    ws["A2"].font = label_font

    if (data.general_note or "").strip():
        ws.merge_cells("A3:E3")
        ws["A3"] = _cell_value(f"Комментарий: {data.general_note.strip()}")
        ws["A3"].font = label_font
        ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
        header_row = 5
    else:
        header_row = 4

    headers = ["№", "Наименование", "Кол", "Доп упак", "Примечание"]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin

    row = header_row + 1
    for section in _section_order(data):
        cells = [c for c in data.cells if c.section == section]
        if not cells:
            continue
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        sec = ws.cell(row=row, column=1, value=_cell_value(section))
        sec.fill = section_fill
        sec.font = section_font
        sec.alignment = Alignment(vertical="center")
        for col in range(1, 6):
            ws.cell(row=row, column=col).border = thin
            ws.cell(row=row, column=col).fill = section_fill
        row += 1

        for item in cells:
            values = [
                "" if item.is_note_only else (item.num if item.num is not None else ""),
                item.name,
                "" if item.is_note_only else item.qty,
                "" if item.is_note_only else item.packing,
                item.note,
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=_cell_value(value))
                cell.border = thin
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            row += 1

    widths = [6, 42, 10, 14, 28]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = f"A{header_row + 1}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_usa_xlsx(data: DismantleMapOut) -> bytes:
    """USA form: Part name RUS/ENG, Qty (1–15), Notes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Карта разбора (США)"

    thin = Border(
        left=Side(style="thin", color="D4D4D4"),
        right=Side(style="thin", color="D4D4D4"),
        top=Side(style="thin", color="D4D4D4"),
        bottom=Side(style="thin", color="D4D4D4"),
    )
    header_fill = PatternFill("solid", fgColor="0D3F10")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    section_fill = PatternFill("solid", fgColor="E8F5E9")
    section_font = Font(bold=True, color="0D3F10", size=11)
    title_font = Font(bold=True, size=14, color="0D3F10")
    label_font = Font(size=10, color="666666")
    yellow = PatternFill("solid", fgColor="FFF59D")

    ws.merge_cells("A1:D1")
    ws["A1"] = _cell_value(data.title or "Карта разбора (США)")
    ws["A1"].font = title_font

    ws.merge_cells("A2:D2")
    hint = data.vehicle_header_hint or "MAKE, MODEL, LOT#, VIN"
    ws["A2"] = _cell_value(f"{hint}: {data.vehicle_label or '—'}")
    ws["A2"].font = label_font

    if (data.general_note or "").strip():
        ws.merge_cells("A3:D3")
        ws["A3"] = _cell_value(f"Notes / Комментарий: {data.general_note.strip()}")
        ws["A3"].font = label_font
        ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
        usa_header_row = 5
    else:
        usa_header_row = 4

    headers = ["Part name (RUS)", "Part name (ENG)", "Снять", "Примечания | Notes"]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=usa_header_row, column=col, value=text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin

    row = usa_header_row + 1
    for section in _section_order(data):
        cells = [c for c in data.cells if c.section == section]
        if not cells:
            continue
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        sec = ws.cell(row=row, column=1, value=_cell_value(section))
        sec.fill = section_fill
        sec.font = section_font
        for col in range(1, 5):
            ws.cell(row=row, column=col).border = thin
            ws.cell(row=row, column=col).fill = section_fill
        row += 1

        for item in cells:
            values = [item.name, item.name_en or "", item.qty, item.note]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=_cell_value(value))
                cell.border = thin
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if col in (3, 4):
                    cell.fill = yellow
            row += 1

    widths = [42, 42, 12, 36]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = f"A{usa_header_row + 1}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
